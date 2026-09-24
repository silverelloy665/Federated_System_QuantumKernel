"""
CI Consistency and Integrity Test Suite for SwarmGuard Modeling Phase.
Verifies that SWARMGUARD_MODELING_LOG.xlsx contains all 8 sheets with real data,
and validates the 12-qubit decoupled QCNN architecture and parameter counts.
"""

import sys
import pytest
import warnings
from pathlib import Path

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
import openpyxl

repo_root = Path(__file__).resolve().parent.parent
if str(repo_root) not in sys.path:
    sys.path.insert(0, str(repo_root))

from swarmguard_pipeline.config import PipelineConfig
from swarmguard_modeling.qcnn_ansatz import QCNNModel

@pytest.fixture
def config():
    return PipelineConfig()

def test_qcnn_ansatz_parameter_budget():
    """Verifies that the 12-qubit QCNN ansatz has exactly 36 trainable parameters."""
    for branch in ["Network", "Physical"]:
        model = QCNNModel(branch, 12)
        meta = model.get_circuit_metadata()
        assert meta["trainable_parameters"] == 36, f"{branch} QCNN parameter count mismatch: {meta['trainable_parameters']} != 36"
        assert meta["num_qubits"] == 12
        assert meta["circuit_depth"] > 0

def test_modeling_log_workbook_structure(config):
    """Verifies that SWARMGUARD_MODELING_LOG.xlsx exists, has all 8 sheets, and is populated."""
    log_path = config.reports_dir / "SWARMGUARD_MODELING_LOG.xlsx"
    if not log_path.exists():
        alt = repo_root / "reports" / "SWARMGUARD_MODELING_LOG.xlsx"
        if alt.exists():
            log_path = alt

    assert log_path.exists(), f"SWARMGUARD_MODELING_LOG.xlsx not found at {log_path}"

    wb = openpyxl.load_workbook(log_path, data_only=True)
    expected_sheets = [
        "Log",
        "0_Angle_Domain_Fix",
        "1_QCNN_Ansatz",
        "2_Control_Matrix",
        "3_Optimizers",
        "4_Federated_Training",
        "5_Alert_Fusion",
        "6_QPU_Submission"
    ]

    assert len(wb.sheetnames) == 8, f"Expected 8 sheets in modeling log, found {len(wb.sheetnames)}: {wb.sheetnames}"
    for s_name in expected_sheets:
        assert s_name in wb.sheetnames, f"Missing sheet '{s_name}' in modeling log"
        ws = wb[s_name]
        assert ws.max_row >= 5, f"Sheet '{s_name}' appears empty (max_row = {ws.max_row})"

